"""
Analytics API — time-series metrics history.
"""

import io
from datetime import datetime, timedelta
from typing import List, Optional
from fastapi import APIRouter, Query
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import StreamingResponse

from app.db.session import get_db_cursor, rows_to_dicts
from app.core.config import get_settings
from app.services.capacity_projection import compute_daily_trend, forecast_capacity


router = APIRouter(prefix="/analytics", tags=["analytics"])
settings = get_settings()
SCHEMA = settings.db_schema



def _fetch_history(array_name: str, hours: int, limit: int) -> List[dict]:
    with get_db_cursor() as cursor:
        cursor.execute(
            f"""SELECT TOP {limit}
                array_name, collected_at,
                read_iops, write_iops,
                read_latency_us, write_latency_us,
                read_bandwidth, write_bandwidth,
                capacity_total, capacity_used, capacity_used_pct, data_reduction
            FROM {SCHEMA}.metrics_history
            WHERE array_name=?
              AND collected_at >= DATEADD(HOUR, -?, GETDATE())
            ORDER BY collected_at ASC""",
            (array_name, hours),
        )
        return rows_to_dicts(cursor, cursor.fetchall())


def _fetch_daily_stats(days: int) -> List[dict]:
    with get_db_cursor() as cursor:
        cursor.execute(
            f"""SELECT TOP {days} * FROM {SCHEMA}.daily_stats
            ORDER BY stat_date DESC""",
        )
        return rows_to_dicts(cursor, cursor.fetchall())


@router.get("/history/{array_name}")
async def get_array_history(
    array_name: str,
    hours: int = Query(default=24, ge=1, le=720),
    limit: int = Query(default=1440, ge=1, le=10000),
):
    """Time-series metrics for a single array (default last 24h)."""
    rows = await run_in_threadpool(_fetch_history, array_name, hours, limit)
    return {"array_name": array_name, "hours": hours, "data_points": len(rows), "data": rows}


@router.get("/daily")
async def get_daily_stats(days: int = Query(default=30, ge=1, le=365)):
    """Daily aggregated stats across all arrays."""
    rows = await run_in_threadpool(_fetch_daily_stats, days)
    return {"days": days, "data": rows}


@router.get("/daily-trend")
async def get_daily_trend(days: int = Query(default=90, ge=1, le=365)):
    """
    Fleet capacity time-series for charting: total usable/used TB and average
    utilization per day, ascending by date. Backed by the daily_stats table.

    Includes a server-computed `projection` (net change, avg rate, headroom,
    projected-full date, trend classification) plus `last_collected` so the UI
    can render a single consistent growth summary and a freshness badge.

    The projection math lives in app/services/capacity_projection.py so it's
    shared with the capacity alerting scheduler job (single source of truth).
    """
    result = await run_in_threadpool(compute_daily_trend, days)
    return {
        "days": days,
        "data_points": len(result["data"]),
        "data": result["data"],
        "last_collected": result["last_collected"],
        "projection": result["projection"],
    }




@router.get("/fleet-history")
async def get_fleet_history(
    hours: int = Query(default=24, ge=1, le=720),
    limit: int = Query(default=5000, ge=1, le=20000),
):
    """Time-series metrics for ALL arrays in a single query."""
    rows = await run_in_threadpool(_fetch_fleet_history, hours, limit)
    arrays = list({r["array_name"] for r in rows})
    return {"hours": hours, "arrays": arrays, "data_points": len(rows), "data": rows}


def _fetch_fleet_history(hours: int, limit: int) -> List[dict]:
    with get_db_cursor() as cursor:
        # Take the TOP N by collected_at DESC (the MOST RECENT rows), then re-sort
        # ASC for charting. The old `TOP N ... ORDER BY ASC` returned the OLDEST N
        # and silently dropped "now": ~90 arrays over 24h is ~26k rows > the 5000
        # cap, so the charts showed only the oldest few hours.
        cursor.execute(
            f"""SELECT array_name, collected_at,
                       read_iops, write_iops,
                       read_latency_us, write_latency_us,
                       read_bandwidth, write_bandwidth,
                       controller_load, queue_depth, nic_util_pct, san_latency_us, queue_latency_us,
                       capacity_used_pct
                FROM (
                    SELECT TOP {limit}
                        array_name, collected_at,
                        read_iops, write_iops,
                        read_latency_us, write_latency_us,
                        read_bandwidth, write_bandwidth,
                        controller_load, queue_depth, nic_util_pct, san_latency_us, queue_latency_us,
                        capacity_used_pct
                    FROM {SCHEMA}.metrics_history
                    WHERE collected_at >= DATEADD(HOUR, -?, GETDATE())
                    ORDER BY collected_at DESC
                ) t
                ORDER BY collected_at ASC""",
            (hours,),
        )
        return rows_to_dicts(cursor, cursor.fetchall())


# ---------------------------------------------------------------------------
# Fleet TREND — server-side time-bucketed aggregate for the interactive
# Analytics chart. GROUP BY a floored time bucket per array (same streaming
# GROUP-BY pattern as capacity-history) so we avoid the two problems the old
# raw fleet-history had for charting:
#   1. The TOP-N row cap silently truncated 24h/7d "all arrays" to the last few
#      hours (~90 arrays * 288 samples/day > cap).
#   2. Each array collected at slightly different timestamps, so a client pivot
#      keyed on exact collected_at produced one-array-per-row (all others NULL)
#      and the lines never connected.
# Bucketing on the server aligns every array onto a shared time grid AND shrinks
# the payload to (arrays * buckets) aggregated rows.
# ---------------------------------------------------------------------------
def _bucket_minutes_for(hours: int) -> int:
    if hours <= 6:
        return 5
    if hours <= 24:
        return 15
    if hours <= 48:
        return 30
    if hours <= 168:
        return 60
    return 180


_TREND_METRIC_COLS = (
    "read_iops", "write_iops",
    "read_latency_us", "write_latency_us",
    "read_bandwidth", "write_bandwidth",
    "controller_load", "nic_util_pct",
    "san_latency_us", "queue_latency_us",
    "capacity_used_pct",
)

_fleet_trend_cache: dict = {}
_FLEET_TREND_TTL = 60


@router.get("/fleet-trend")
async def get_fleet_trend(
    hours: int = Query(default=24, ge=1, le=720),
    bucket_min: int = Query(default=0, ge=0, le=1440),
):
    """Time-bucketed fleet metrics for the interactive Analytics chart. One AVG
    aggregate per array per bucket, across every metric family. Cached 60s per
    (hours, bucket)."""
    bucket = bucket_min or _bucket_minutes_for(hours)
    key = (hours, bucket)
    hit = _fleet_trend_cache.get(key)
    if hit and (_time.time() - hit[0]) < _FLEET_TREND_TTL:
        return hit[1]
    rows = await run_in_threadpool(_fetch_fleet_trend, hours, bucket)
    arrays = sorted({r["array_name"] for r in rows})
    payload = {
        "hours": hours,
        "bucket_min": bucket,
        "arrays": arrays,
        "metrics": list(_TREND_METRIC_COLS),
        "data_points": len(rows),
        "data": rows,
    }
    _fleet_trend_cache[key] = (_time.time(), payload)
    return payload


def _fetch_fleet_trend(hours: int, bucket: int) -> List[dict]:
    # Floor collected_at to `bucket`-minute boundaries: minutes-since-datetime0
    # integer-divided by bucket, times bucket, added back onto datetime 0.
    #
    # `bucket` is INLINED as an int literal (not a parameter). SQL Server matches
    # the SELECT bucket expression to the GROUP BY one textually; two positional
    # `?` markers aren't considered the same expression, so a parameterized bucket
    # trips "collected_at is invalid in the select list". `bucket` is a bounded
    # int from a range-checked Query / our own table, so inlining is injection-safe.
    bucket = max(1, int(bucket))
    bucket_expr = (
        f"DATEADD(MINUTE, (DATEDIFF(MINUTE, 0, collected_at) / {bucket}) * {bucket}, 0)"
    )
    avg_cols = ", ".join(
        f"AVG(CAST({c} AS FLOAT)) AS {c}" for c in _TREND_METRIC_COLS
    )
    with get_db_cursor() as cursor:
        cursor.execute(
            f"""SELECT array_name,
                       {bucket_expr} AS bucket,
                       {avg_cols}
                FROM {SCHEMA}.metrics_history WITH (NOLOCK)
                WHERE collected_at >= DATEADD(HOUR, -?, GETDATE())
                GROUP BY array_name, {bucket_expr}
                ORDER BY bucket ASC""",
            (hours,),
        )
        return rows_to_dicts(cursor, cursor.fetchall())


# Small in-memory TTL cache: the Capacity page auto-fetches this and react-query
# refetches on focus/remount across users, so without a cache the (heavy) query
# runs many times concurrently. Cache per `days` for 5 min.
import time as _time
_caphist_cache: dict = {}
_CAPHIST_TTL = 300


@router.get("/capacity-history")
async def get_capacity_history(days: int = Query(default=90, ge=1, le=400)):
    """Per-array DAILY capacity history + array metadata, for the slicer-driven
    capacity view. One fetch; the frontend filters (vendor/group/model) and
    re-aggregates the trends client-side. Excludes disabled arrays. Cached 5 min."""
    hit = _caphist_cache.get(days)
    if hit and (_time.monotonic() - hit[0]) < _CAPHIST_TTL:
        return hit[1]
    result = await run_in_threadpool(_fetch_capacity_history, days)
    _caphist_cache[days] = (_time.monotonic(), result)
    return result


def _fetch_capacity_history(days: int) -> dict:
    with get_db_cursor() as cursor:
        # array metadata (vendor / group / model), disabled arrays excluded, joined
        # to metrics_current so we only list arrays we actually collect capacity for.
        cursor.execute(
            f"""SELECT mc.array_name,
                       COALESCE(ma.vendor, mc.vendor, 'unknown')          AS vendor,
                       COALESCE(NULLIF(ma.group_label, ''), 'Unassigned') AS grp,
                       COALESCE(ma.model, '')                             AS model
                FROM {SCHEMA}.metrics_current mc WITH (NOLOCK)
                LEFT JOIN {SCHEMA}.managed_arrays ma WITH (NOLOCK) ON ma.array_name = mc.array_name
                WHERE mc.array_name NOT IN (
                    SELECT array_name FROM {SCHEMA}.managed_arrays WITH (NOLOCK) WHERE enabled=0)"""
        )
        arrays = [
            {"array_name": r[0], "vendor": r[1], "group": r[2], "model": r[3]}
            for r in cursor.fetchall()
        ]
        # One daily point per array via a GROUP BY aggregate (AVG over the day —
        # capacity moves slowly, so the daily mean is a faithful trend value). This
        # replaced a ROW_NUMBER window over the full ~1.3M-row 90-day set, which
        # cost ~5s and — fired concurrently by the auto-loading Capacity page —
        # blew past the query timeout and wedged the worker. GROUP BY streams the
        # scan straight to ~one row per array/day with no giant sorted intermediate.
        cursor.execute(
            f"""SELECT array_name, CONVERT(varchar(10), CAST(collected_at AS DATE), 23) AS day,
                       AVG(CAST(capacity_used AS FLOAT))  AS used,
                       AVG(CAST(capacity_total AS FLOAT)) AS total,
                       AVG(capacity_used_pct)             AS pct
                FROM {SCHEMA}.metrics_history WITH (NOLOCK)
                WHERE collected_at >= DATEADD(DAY, -?, GETDATE())
                  AND capacity_total > 0
                GROUP BY array_name, CAST(collected_at AS DATE)
                ORDER BY day, array_name""",
            (days,),
        )
        points = []
        for arr, day, used, total, pct in cursor.fetchall():
            points.append({
                "d": day,
                "a": arr,
                "used_tb": round((used or 0) / _TIB, 2),
                "total_tb": round((total or 0) / _TIB, 2),
                "used_pct": round(pct or 0, 2),
            })
    return {"days": days, "arrays": arrays, "points": points}


# ── Capacity breakdown (by vendor / cloud / vendor×cloud) ─────────────────────

_TIB = 1099511627776.0  # bytes per TiB


def _fetch_capacity_breakdown() -> dict:
    """
    Aggregate usable/used/free capacity across the fleet, grouped by
    vendor, by cloud/group, and by the vendor×cloud pivot.

    Joins live metrics (metrics_current) with the managed_arrays inventory
    so each array is attributed to its vendor and its cloud/site group.
    """
    # LEFT JOIN so arrays present in metrics_current but missing from
    # managed_arrays still appear (group falls back to 'Unassigned').
    base_sql = f"""
        SELECT
            mc.array_name,
            COALESCE(ma.vendor, mc.vendor, 'unknown')              AS vendor,
            COALESCE(NULLIF(ma.group_label, ''), 'Unassigned')     AS cloud,
            CAST(mc.capacity_total AS FLOAT)                       AS capacity_total,
            CAST(mc.capacity_used  AS FLOAT)                       AS capacity_used
        FROM {SCHEMA}.metrics_current mc WITH (NOLOCK)
        LEFT JOIN {SCHEMA}.managed_arrays ma WITH (NOLOCK)
            ON ma.array_name = mc.array_name
        -- Exclude arrays explicitly disabled in managed_arrays so the Capacity
        -- page fleet totals match fleet-stats (which already excludes them);
        -- arrays absent from managed_arrays (enabled IS NULL) still count.
        WHERE ISNULL(ma.enabled, 1) = 1
    """

    with get_db_cursor() as cursor:
        cursor.execute(base_sql)
        rows = rows_to_dicts(cursor, cursor.fetchall())

    def _blank() -> dict:
        return {"arrays": 0, "usable_bytes": 0.0, "used_bytes": 0.0}

    by_vendor: dict = {}
    by_cloud: dict = {}
    pivot: dict = {}  # key = "vendor||cloud"
    fleet = _blank()

    for r in rows:
        vendor = (r.get("vendor") or "unknown").lower()
        cloud = r.get("cloud") or "Unassigned"
        total = float(r.get("capacity_total") or 0)
        used = float(r.get("capacity_used") or 0)

        for bucket, key in ((by_vendor, vendor), (by_cloud, cloud)):
            b = bucket.setdefault(key, _blank())
            b["arrays"] += 1
            b["usable_bytes"] += total
            b["used_bytes"] += used

        pk = f"{vendor}||{cloud}"
        p = pivot.setdefault(pk, {"vendor": vendor, "cloud": cloud, **_blank()})
        p["arrays"] += 1
        p["usable_bytes"] += total
        p["used_bytes"] += used

        fleet["arrays"] += 1
        fleet["usable_bytes"] += total
        fleet["used_bytes"] += used

    def _finalize(b: dict) -> dict:
        usable = b["usable_bytes"]
        used = b["used_bytes"]
        free = max(usable - used, 0.0)
        return {
            "arrays": b["arrays"],
            "usable_tb": round(usable / _TIB, 2),
            "used_tb": round(used / _TIB, 2),
            "free_tb": round(free / _TIB, 2),
            "allocated_tb": round(usable / _TIB, 2),
            "utilization_pct": round((used / usable * 100), 1) if usable > 0 else 0.0,
        }

    return {
        "fleet": _finalize(fleet),
        "by_vendor": [
            {"vendor": k, **_finalize(v)}
            for k, v in sorted(by_vendor.items(), key=lambda x: -x[1]["usable_bytes"])
        ],
        "by_cloud": [
            {"cloud": k, **_finalize(v)}
            for k, v in sorted(by_cloud.items(), key=lambda x: -x[1]["usable_bytes"])
        ],
        "by_vendor_cloud": [
            {"vendor": v["vendor"], "cloud": v["cloud"], **_finalize(v)}
            for v in sorted(pivot.values(), key=lambda x: -x["usable_bytes"])
        ],
    }


@router.get("/capacity-breakdown")
async def get_capacity_breakdown():
    """
    Usable / used / allocated capacity broken down by platform (vendor),
    by cloud (group_label), and as a vendor×cloud pivot. All sizes in TB.
    """
    return await run_in_threadpool(_fetch_capacity_breakdown)


# ── Per-array YTD / trailing growth ───────────────────────────────────────────

def _fetch_array_growth(array_name: str, months: int) -> dict:
    """
    Capacity growth for a single array:
      - YTD growth (earliest sample on/after Jan 1 of current year vs latest)
      - Trailing-N-month trend (monthly capacity_used / capacity_total samples)

    Reads from metrics_history. Requires history retention >= the requested
    window (see ANALYTICS_HISTORY_RETENTION_DAYS / cleanup_old_history).
    """
    with get_db_cursor() as cursor:
        # Monthly trend: first sample of each month within the window.
        cursor.execute(
            f"""
            WITH ranked AS (
                SELECT
                    collected_at,
                    capacity_total,
                    capacity_used,
                    capacity_used_pct,
                    ROW_NUMBER() OVER (
                        PARTITION BY YEAR(collected_at), MONTH(collected_at)
                        ORDER BY collected_at ASC
                    ) AS rn
                FROM {SCHEMA}.metrics_history WITH (NOLOCK)
                WHERE array_name = ?
                  AND collected_at >= DATEADD(MONTH, -?, GETDATE())
            )
            SELECT collected_at, capacity_total, capacity_used, capacity_used_pct
            FROM ranked
            WHERE rn = 1
            ORDER BY collected_at ASC
            """,
            (array_name, months),
        )
        trend = rows_to_dicts(cursor, cursor.fetchall())

        # Earliest YTD sample (on/after Jan 1 this year).
        cursor.execute(
            f"""
            SELECT TOP 1 collected_at, capacity_total, capacity_used
            FROM {SCHEMA}.metrics_history WITH (NOLOCK)
            WHERE array_name = ?
              AND collected_at >= DATEFROMPARTS(YEAR(GETDATE()), 1, 1)
            ORDER BY collected_at ASC
            """,
            (array_name,),
        )
        ytd_start = row = cursor.fetchone()
        ytd_start = dict(zip([c[0] for c in cursor.description], row)) if row else None

    # Current values come from metrics_current (always present, freshest).
    with get_db_cursor() as cursor:
        cursor.execute(
            f"""SELECT capacity_total, capacity_used, capacity_used_pct, collected_at
                FROM {SCHEMA}.metrics_current WITH (NOLOCK) WHERE array_name = ?""",
            (array_name,),
        )
        crow = cursor.fetchone()
        current = (
            dict(zip([c[0] for c in cursor.description], crow)) if crow else None
        )

    def _to_tb(v) -> float:
        return round(float(v or 0) / _TIB, 2)

    ytd = None
    if ytd_start and current:
        start_used = float(ytd_start.get("capacity_used") or 0)
        cur_used = float(current.get("capacity_used") or 0)
        delta = cur_used - start_used
        ytd = {
            "start_date": str(ytd_start.get("collected_at")),
            "start_used_tb": _to_tb(start_used),
            "current_used_tb": _to_tb(cur_used),
            "growth_tb": round(delta / _TIB, 2),
            "growth_pct": round((delta / start_used * 100), 1) if start_used > 0 else None,
        }

    return {
        "array_name": array_name,
        "months": months,
        "current": {
            "usable_tb": _to_tb(current.get("capacity_total")) if current else None,
            "used_tb": _to_tb(current.get("capacity_used")) if current else None,
            "used_pct": current.get("capacity_used_pct") if current else None,
            "collected_at": str(current.get("collected_at")) if current else None,
        },
        "ytd": ytd,
        "trend": [
            {
                "date": str(t.get("collected_at")),
                "usable_tb": _to_tb(t.get("capacity_total")),
                "used_tb": _to_tb(t.get("capacity_used")),
                "used_pct": t.get("capacity_used_pct"),
            }
            for t in trend
        ],
    }


@router.get("/array-growth/{array_name}")
async def get_array_growth(
    array_name: str,
    months: int = Query(default=12, ge=1, le=24),
):
    """
    Per-array capacity growth: YTD growth + trailing-N-month monthly trend.
    Depends on metrics_history retention (default extended to 365 days).
    """
    return await run_in_threadpool(_fetch_array_growth, array_name, months)


@router.get("/forecast")
async def get_forecast(
    array: Optional[str] = Query(default=None, description="array_name; omit for fleet-wide"),
    target_date: Optional[str] = Query(default=None, description="YYYY-MM-DD"),
    window: int = Query(default=90, ge=7, le=365, description="trailing days to fit the trend"),
):
    """Project used capacity forward for an array (or the whole fleet) to an
    optional target date, using its historical trend. Deterministic; the chat
    narrates this result rather than inventing numbers."""
    return await run_in_threadpool(forecast_capacity, array, target_date, window)


def _fetch_top_growers(days: int, limit: int) -> List[dict]:
    """
    Rank every array by capacity-used delta over the trailing window.

    For each array we compare the earliest sample within the window
    (from metrics_history) to its current value (metrics_current),
    yielding growth (positive) or decline (negative) in TB and %.

    Single batched query using ROW_NUMBER so we don't fan out per array.
    """
    with get_db_cursor() as cursor:
        # Earliest sample per array within the window.
        cursor.execute(
            f"""
            WITH ranked AS (
                SELECT
                    array_name,
                    capacity_used,
                    capacity_total,
                    collected_at,
                    ROW_NUMBER() OVER (
                        PARTITION BY array_name
                        ORDER BY collected_at ASC
                    ) AS rn
                FROM {SCHEMA}.metrics_history WITH (NOLOCK)
                WHERE collected_at >= DATEADD(DAY, -?, GETDATE())
            )
            SELECT array_name, capacity_used, capacity_total, collected_at
            FROM ranked
            WHERE rn = 1
            """,
            (days,),
        )
        start_rows = rows_to_dicts(cursor, cursor.fetchall())

    start_by_array = {r["array_name"]: r for r in start_rows}

    with get_db_cursor() as cursor:
        cursor.execute(
            f"""
            SELECT
                mc.array_name,
                COALESCE(ma.vendor, mc.vendor, 'unknown') AS vendor,
                CAST(mc.capacity_total AS FLOAT)          AS capacity_total,
                CAST(mc.capacity_used  AS FLOAT)          AS capacity_used,
                mc.capacity_used_pct
            FROM {SCHEMA}.metrics_current mc WITH (NOLOCK)
            LEFT JOIN {SCHEMA}.managed_arrays ma WITH (NOLOCK)
                ON ma.array_name = mc.array_name
            WHERE ISNULL(ma.enabled, 1) = 1
            """,
        )
        current_rows = rows_to_dicts(cursor, cursor.fetchall())

    out: List[dict] = []
    for cur in current_rows:
        name = cur["array_name"]
        start = start_by_array.get(name)
        if not start:
            continue  # no history sample in window — skip
        start_used = float(start.get("capacity_used") or 0)
        cur_used = float(cur.get("capacity_used") or 0)
        delta = cur_used - start_used
        total = float(cur.get("capacity_total") or 0)
        out.append({
            "array_name": name,
            "vendor": (cur.get("vendor") or "unknown").lower(),
            "start_date": str(start.get("collected_at")),
            "start_used_tb": round(start_used / _TIB, 2),
            "current_used_tb": round(cur_used / _TIB, 2),
            "delta_tb": round(delta / _TIB, 2),
            "delta_pct": round((delta / start_used * 100), 1) if start_used > 0 else None,
            "utilization_pct": cur.get("capacity_used_pct"),
        })

    # Sort by absolute change so both biggest growers and biggest decliners
    # bubble to the top; the frontend can split positive/negative.
    out.sort(key=lambda x: -abs(x["delta_tb"]))
    return out[:limit]


@router.get("/top-growers")
async def get_top_growers(
    days: int = Query(default=90, ge=1, le=365),
    limit: int = Query(default=20, ge=1, le=200),
):
    """
    Arrays ranked by capacity-used change over the trailing window.
    Positive delta = growth, negative = decline. Sorted by magnitude so the
    UI can show both top growers and top shrinkers.
    """
    rows = await run_in_threadpool(_fetch_top_growers, days, limit)
    return {"days": days, "count": len(rows), "data": rows}



# ── Per-volume growth (backed by volumes_history) ─────────────────────────────

def _fetch_volume_history_coverage() -> dict:
    """
    How much per-volume time-series we actually have: earliest/latest snapshot,
    distinct snapshot timestamps, and row count. Lets the UI tell the user the
    real data window before drawing growth charts.
    """
    with get_db_cursor() as cursor:
        cursor.execute(
            f"""
            SELECT
                MIN(collected_at) AS first_seen,
                MAX(collected_at) AS last_seen,
                COUNT(*)          AS rows_total,
                COUNT(DISTINCT CAST(collected_at AS DATE)) AS distinct_days
            FROM {SCHEMA}.volumes_history WITH (NOLOCK)
            """,
        )
        row = cursor.fetchone()
        cov = dict(zip([c[0] for c in cursor.description], row)) if row else {}

    return {
        "first_seen": str(cov.get("first_seen")) if cov.get("first_seen") else None,
        "last_seen": str(cov.get("last_seen")) if cov.get("last_seen") else None,
        "rows_total": int(cov.get("rows_total") or 0),
        "distinct_days": int(cov.get("distinct_days") or 0),
    }


@router.get("/volume-history-coverage")
async def get_volume_history_coverage():
    """
    Report the available per-volume history window (first/last snapshot,
    distinct days, total rows). Useful before rendering volume-growth views,
    since volumes_history only starts accumulating once collectors snapshot.
    """
    return await run_in_threadpool(_fetch_volume_history_coverage)


def _fetch_volume_growth(array_name: str, volume_name: str, days: int) -> dict:
    """
    Per-volume used/size trend over the trailing window.

    One sample per snapshot timestamp from volumes_history (already one row per
    volume per collection), ascending by time for left-to-right charting.
    """
    with get_db_cursor() as cursor:
        cursor.execute(
            f"""
            SELECT collected_at, size, used, data_reduction, snapshots
            FROM {SCHEMA}.volumes_history WITH (NOLOCK)
            WHERE array_name = ? AND volume_name = ?
              AND collected_at >= DATEADD(DAY, -?, GETDATE())
            ORDER BY collected_at ASC
            """,
            (array_name, volume_name, days),
        )
        trend = rows_to_dicts(cursor, cursor.fetchall())

    series = [
        {
            "date": str(t.get("collected_at")),
            "size_tb": round(float(t.get("size") or 0) / _TIB, 3),
            "used_tb": round(float(t.get("used") or 0) / _TIB, 3),
            "data_reduction": round(float(t.get("data_reduction") or 1), 2),
            "snapshots": int(t.get("snapshots") or 0),
        }
        for t in trend
    ]

    growth = None
    if len(series) >= 2:
        start_used = series[0]["used_tb"]
        cur_used = series[-1]["used_tb"]
        delta = round(cur_used - start_used, 3)
        growth = {
            "start_date": series[0]["date"],
            "start_used_tb": start_used,
            "current_used_tb": cur_used,
            "growth_tb": delta,
            "growth_pct": round((delta / start_used * 100), 1) if start_used > 0 else None,
        }

    return {
        "array_name": array_name,
        "volume_name": volume_name,
        "days": days,
        "data_points": len(series),
        "growth": growth,
        "trend": series,
    }


@router.get("/volume-growth/{array_name}/{volume_name:path}")
async def get_volume_growth(
    array_name: str,
    volume_name: str,
    days: int = Query(default=90, ge=1, le=365),
):
    """
    Per-volume capacity growth/decline trend over the trailing window.
    Backed by volumes_history (one snapshot per collection). volume_name is a
    path param so vendor names containing '/' (Oracle, NetApp svm:vol) work.
    """
    return await run_in_threadpool(_fetch_volume_growth, array_name, volume_name, days)


def _fetch_top_volume_growers(days: int, limit: int, array_name: Optional[str]) -> List[dict]:
    """
    Rank volumes by used-capacity delta over the trailing window.

    For each (array, volume) we take the earliest snapshot inside the window
    (ROW_NUMBER) and join it to the current value in volumes_cache. Positive
    delta = growth, negative = decline. Single batched query — no per-volume fan-out.
    """
    array_filter = " AND array_name = ?" if array_name else ""
    hist_params = [days] + ([array_name] if array_name else [])

    with get_db_cursor() as cursor:
        cursor.execute(
            f"""
            WITH ranked AS (
                SELECT
                    array_name, vendor, volume_name, used, size, collected_at,
                    ROW_NUMBER() OVER (
                        PARTITION BY array_name, volume_name
                        ORDER BY collected_at ASC
                    ) AS rn
                FROM {SCHEMA}.volumes_history WITH (NOLOCK)
                WHERE collected_at >= DATEADD(DAY, -?, GETDATE())
                {array_filter}
            )
            SELECT array_name, vendor, volume_name, used, size, collected_at
            FROM ranked
            WHERE rn = 1
            """,
            tuple(hist_params),
        )
        start_rows = rows_to_dicts(cursor, cursor.fetchall())

    start_by_key = {(r["array_name"], r["volume_name"]): r for r in start_rows}

    cur_filter = " WHERE array_name = ?" if array_name else ""
    cur_params = (array_name,) if array_name else ()
    with get_db_cursor() as cursor:
        cursor.execute(
            f"""
            SELECT array_name, vendor, volume_name, size, used
            FROM {SCHEMA}.volumes_cache WITH (NOLOCK)
            {cur_filter}
            """,
            cur_params,
        )
        current_rows = rows_to_dicts(cursor, cursor.fetchall())

    out: List[dict] = []
    for cur in current_rows:
        key = (cur["array_name"], cur["volume_name"])
        start = start_by_key.get(key)
        if not start:
            continue  # no history sample in window — skip
        start_used = float(start.get("used") or 0)
        cur_used = float(cur.get("used") or 0)
        delta = cur_used - start_used
        out.append({
            "array_name": cur["array_name"],
            "vendor": (cur.get("vendor") or "unknown").lower(),
            "volume_name": cur["volume_name"],
            "start_date": str(start.get("collected_at")),
            "start_used_tb": round(start_used / _TIB, 3),
            "current_used_tb": round(cur_used / _TIB, 3),
            "delta_tb": round(delta / _TIB, 3),
            "delta_pct": round((delta / start_used * 100), 1) if start_used > 0 else None,
        })

    # Sort by absolute change so biggest growers AND decliners surface together.
    out.sort(key=lambda x: -abs(x["delta_tb"]))
    return out[:limit]


@router.get("/top-volume-growers")
async def get_top_volume_growers(
    days: int = Query(default=90, ge=1, le=365),
    limit: int = Query(default=25, ge=1, le=500),
    array_name: Optional[str] = Query(default=None),
):
    """
    Volumes ranked by used-capacity change over the trailing window.
    Optionally scoped to a single array. Sorted by magnitude so the UI can
    split growth vs decline. Backed by volumes_history + volumes_cache.
    """
    rows = await run_in_threadpool(_fetch_top_volume_growers, days, limit, array_name)
    return {"days": days, "array_name": array_name, "count": len(rows), "data": rows}


# ── Excel capacity export ─────────────────────────────────────────────────────


# Keywords used to classify an array's cloud placement from its
# group_label / site / technology inventory fields.
_CSP_KEYWORDS = {
    "Azure": ("azure", "cbs azure", "msft", "microsoft"),
    "AWS":   ("aws", "amazon", "cbs aws", "ec2"),
    "GCP":   ("gcp", "google"),
    "OCI":   ("oci", "oracle cloud"),
}
_ONPREM_KEYWORDS = ("on-prem", "onprem", "on prem", "datacenter", "data center", "dc", "colo")


def _classify_cloud(group: str, site: str, technology: str) -> tuple[str, str]:
    """
    Returns (cloud_type, csp) where:
      cloud_type ∈ {"Cloud", "On-Prem", "Unknown"}
      csp        ∈ {"Azure", "AWS", "GCP", "OCI", "On-Prem", "Unknown"}
    Classification is keyword-based across group_label, site and technology.
    """
    blob = " ".join(str(x or "").lower() for x in (group, site, technology))

    for csp, kws in _CSP_KEYWORDS.items():
        if any(k in blob for k in kws):
            return "Cloud", csp

    if any(k in blob for k in _ONPREM_KEYWORDS):
        return "On-Prem", "On-Prem"

    # 'technology' column sometimes literally says Cloud / On-Prem / Hybrid
    if "cloud" in blob:
        return "Cloud", "Unknown"
    if "hybrid" in blob:
        return "Cloud", "Hybrid"

    return "Unknown", "Unknown"


def _fetch_array_capacity_rows() -> List[dict]:
    """
    One row per array combining live capacity (metrics_current) with inventory
    (managed_arrays), enriched with derived cloud_type / csp classification.
    """
    sql = f"""
        SELECT
            mc.array_name,
            COALESCE(ma.vendor, mc.vendor, 'unknown')          AS vendor,
            COALESCE(NULLIF(ma.group_label, ''), 'Unassigned') AS group_label,
            ma.site,
            ma.technology,
            ma.model,
            mc.array_status,
            CAST(mc.capacity_total AS FLOAT)                   AS capacity_total,
            CAST(mc.capacity_used  AS FLOAT)                   AS capacity_used,
            mc.capacity_used_pct,
            mc.data_reduction,
            mc.collected_at
        FROM {SCHEMA}.metrics_current mc WITH (NOLOCK)
        LEFT JOIN {SCHEMA}.managed_arrays ma WITH (NOLOCK)
            ON ma.array_name = mc.array_name
        WHERE ISNULL(ma.enabled, 1) = 1
        ORDER BY mc.array_name
    """
    with get_db_cursor() as cursor:
        cursor.execute(sql)
        rows = rows_to_dicts(cursor, cursor.fetchall())

    out: List[dict] = []
    for r in rows:
        total = float(r.get("capacity_total") or 0)
        used = float(r.get("capacity_used") or 0)
        free = max(total - used, 0.0)
        cloud_type, csp = _classify_cloud(
            r.get("group_label"), r.get("site"), r.get("technology")
        )
        out.append({
            "array_name": r.get("array_name"),
            "vendor": (r.get("vendor") or "unknown").lower(),
            "cloud_type": cloud_type,
            "csp": csp,
            "group_label": r.get("group_label"),
            "site": r.get("site") or "",
            "model": r.get("model") or "",
            "status": r.get("array_status") or "",
            "usable_tb": round(total / _TIB, 2),
            "used_tb": round(used / _TIB, 2),
            "free_tb": round(free / _TIB, 2),
            "utilization_pct": round((used / total * 100), 1) if total > 0 else 0.0,
            "data_reduction": round(float(r.get("data_reduction") or 0), 2),
            "collected_at": str(r.get("collected_at") or ""),
        })
    return out


def _summarize(rows: List[dict], key: str) -> List[dict]:
    """Group rows by a key, summing capacity. Returns sorted list of dicts."""
    agg: dict = {}
    for r in rows:
        k = r.get(key) or "Unknown"
        a = agg.setdefault(k, {key: k, "arrays": 0, "usable_tb": 0.0, "used_tb": 0.0, "free_tb": 0.0})
        a["arrays"] += 1
        a["usable_tb"] += r["usable_tb"]
        a["used_tb"] += r["used_tb"]
        a["free_tb"] += r["free_tb"]
    for a in agg.values():
        a["usable_tb"] = round(a["usable_tb"], 2)
        a["used_tb"] = round(a["used_tb"], 2)
        a["free_tb"] = round(a["free_tb"], 2)
        a["utilization_pct"] = round((a["used_tb"] / a["usable_tb"] * 100), 1) if a["usable_tb"] > 0 else 0.0
    return sorted(agg.values(), key=lambda x: -x["usable_tb"])


def _build_capacity_workbook() -> bytes:
    """Build a multi-sheet .xlsx capacity report and return the bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    rows = _fetch_array_capacity_rows()

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=13)

    wb = Workbook()

    def _write_sheet(ws, columns: List[tuple], data: List[dict], title: str):
        ws.cell(row=1, column=1, value=title).font = title_font
        # header row (row 3)
        for ci, (field, label, _w) in enumerate(columns, start=1):
            c = ws.cell(row=3, column=ci, value=label)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center")
        # data rows
        for ri, row in enumerate(data, start=4):
            for ci, (field, _label, _w) in enumerate(columns, start=1):
                ws.cell(row=ri, column=ci, value=row.get(field))
        # column widths + freeze header
        for ci, (_f, _l, w) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A4"

    # Sheet 1: per-array detail
    ws1 = wb.active
    ws1.title = "Arrays"
    _write_sheet(ws1, [
        ("array_name", "Array", 26),
        ("vendor", "Vendor", 12),
        ("cloud_type", "Cloud Type", 12),
        ("csp", "CSP", 12),
        ("group_label", "Group", 20),
        ("site", "Site", 18),
        ("model", "Model", 18),
        ("status", "Status", 12),
        ("usable_tb", "Usable (TB)", 13),
        ("used_tb", "Used (TB)", 12),
        ("free_tb", "Free (TB)", 12),
        ("utilization_pct", "Util %", 9),
        ("data_reduction", "Data Reduction", 14),
        ("collected_at", "Collected At", 22),
    ], rows, f"USM Capacity — Per Array  ({len(rows)} arrays)")

    summary_cols_base = [
        ("arrays", "Arrays", 10),
        ("usable_tb", "Usable (TB)", 13),
        ("used_tb", "Used (TB)", 12),
        ("free_tb", "Free (TB)", 12),
        ("utilization_pct", "Util %", 9),
    ]

    # Sheet 2: by cloud type (On-Prem vs Cloud)
    _write_sheet(
        wb.create_sheet("By Cloud Type"),
        [("cloud_type", "Cloud Type", 16)] + summary_cols_base,
        _summarize(rows, "cloud_type"),
        "Capacity by Cloud Type (On-Prem vs Cloud)",
    )

    # Sheet 3: by CSP
    _write_sheet(
        wb.create_sheet("By CSP"),
        [("csp", "CSP", 16)] + summary_cols_base,
        _summarize(rows, "csp"),
        "Capacity by Cloud Provider (CSP)",
    )

    # Sheet 4: by vendor
    _write_sheet(
        wb.create_sheet("By Vendor"),
        [("vendor", "Vendor", 16)] + summary_cols_base,
        _summarize(rows, "vendor"),
        "Capacity by Platform (Vendor)",
    )

    # Sheet 5: by group
    _write_sheet(
        wb.create_sheet("By Group"),
        [("group_label", "Group", 24)] + summary_cols_base,
        _summarize(rows, "group_label"),
        "Capacity by Group / Cloud Label",
    )

    # Sheet 6: vendor × CSP pivot
    pivot_rows: dict = {}
    for r in rows:
        k = (r["vendor"], r["csp"])
        a = pivot_rows.setdefault(k, {
            "vendor": r["vendor"], "csp": r["csp"],
            "arrays": 0, "usable_tb": 0.0, "used_tb": 0.0, "free_tb": 0.0,
        })
        a["arrays"] += 1
        a["usable_tb"] += r["usable_tb"]
        a["used_tb"] += r["used_tb"]
        a["free_tb"] += r["free_tb"]
    for a in pivot_rows.values():
        a["usable_tb"] = round(a["usable_tb"], 2)
        a["used_tb"] = round(a["used_tb"], 2)
        a["free_tb"] = round(a["free_tb"], 2)
        a["utilization_pct"] = round((a["used_tb"] / a["usable_tb"] * 100), 1) if a["usable_tb"] > 0 else 0.0
    _write_sheet(
        wb.create_sheet("Vendor x CSP"),
        [("vendor", "Vendor", 16), ("csp", "CSP", 14)] + summary_cols_base,
        sorted(pivot_rows.values(), key=lambda x: -x["usable_tb"]),
        "Capacity by Vendor × CSP",
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


@router.get("/capacity-export.xlsx")
async def export_capacity_xlsx():
    """
    Download a multi-sheet Excel capacity report: per-array detail plus
    summaries by cloud type (On-Prem vs Cloud), CSP, vendor, group, and
    a vendor×CSP pivot.
    """
    data = await run_in_threadpool(_build_capacity_workbook)
    fname = f"usm_capacity_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


