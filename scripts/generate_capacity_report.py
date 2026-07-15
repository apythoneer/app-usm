#!/usr/bin/env python3
"""
Standalone capacity-report generator.

Pulls the live USM array list from the ops API and writes a multi-sheet
Excel workbook (per-array detail + grouped summaries by cloud type, CSP,
vendor, group, and a vendor×CSP pivot).

This mirrors the backend endpoint /api/v1/analytics/capacity-export.xlsx but
can be run from a workstation without redeploying the backend.

Usage:
    pip install requests openpyxl
    python scripts/generate_capacity_report.py [--base-url http://host:8000] [--out file.xlsx]
"""
from __future__ import annotations

import argparse
import datetime as dt
import sys

TIB = 1099511627776.0

# Datacenter codes that mean on-prem (3-letter *DC groups).
_ONPREM_CODES = {"ddc", "odc", "idc", "adc", "mdc", "cdc"}
_CSP_KEYWORDS = {
    "Azure": ("azure", "azu", "msft", "microsoft"),
    "AWS":   ("aws", "amazon", "ec2"),
    "GCP":   ("gcp", "google"),
    "OCI":   ("oci", "oracle cloud"),
}


def classify(group: str) -> tuple[str, str]:
    """Return (cloud_type, csp) from an array's group label."""
    g = (group or "").strip().lower()
    if g in _ONPREM_CODES:
        return "On-Prem", "On-Prem"
    if g.startswith("cloud") or "cloud" in g:
        for csp, kws in _CSP_KEYWORDS.items():
            if any(k in g for k in kws):
                return "Cloud", csp
        return "Cloud", "Unknown"
    for csp, kws in _CSP_KEYWORDS.items():
        if any(k in g for k in kws):
            return "Cloud", csp
    # bare datacenter-style code we didn't list -> treat as on-prem
    if len(g) == 3 and g.endswith("dc"):
        return "On-Prem", "On-Prem"
    return "Unknown", "Unknown"


def fetch_arrays(base_url: str) -> list[dict]:
    import requests
    url = base_url.rstrip("/") + "/api/v1/arrays"
    r = requests.get(url, timeout=30)
    r.raise_for_status()
    return r.json()


def to_rows(arrays: list[dict]) -> list[dict]:
    out = []
    for a in arrays:
        total = float(a.get("capacity_total_bytes") or 0)
        used_pct = float(a.get("capacity_used_pct") or 0)
        used = total * used_pct / 100.0
        free = max(total - used, 0.0)
        cloud_type, csp = classify(a.get("group", ""))
        out.append({
            "array_name": a.get("array_name"),
            "vendor": (a.get("vendor") or "unknown").lower(),
            "cloud_type": cloud_type,
            "csp": csp,
            "group": a.get("group") or "",
            "model": a.get("model") or "",
            "status": a.get("array_status") or "",
            "usable_tb": round(total / TIB, 2),
            "used_tb": round(used / TIB, 2),
            "free_tb": round(free / TIB, 2),
            "utilization_pct": round(used_pct, 1),
            "data_reduction": round(float(a.get("data_reduction") or 0), 2),
            "collected_at": str(a.get("collected_at") or ""),
        })
    return sorted(out, key=lambda x: (x["vendor"], x["array_name"]))


def summarize(rows: list[dict], key: str) -> list[dict]:
    agg: dict = {}
    for r in rows:
        k = r.get(key) or "Unknown"
        a = agg.setdefault(k, {key: k, "arrays": 0, "usable_tb": 0.0, "used_tb": 0.0, "free_tb": 0.0})
        a["arrays"] += 1
        a["usable_tb"] += r["usable_tb"]
        a["used_tb"] += r["used_tb"]
        a["free_tb"] += r["free_tb"]
    for a in agg.values():
        a["usable_tb"] = round(a["usable_tb"], 2)
        a["used_tb"] = round(a["used_tb"], 2)
        a["free_tb"] = round(a["free_tb"], 2)
        a["utilization_pct"] = round(a["used_tb"] / a["usable_tb"] * 100, 1) if a["usable_tb"] > 0 else 0.0
    return sorted(agg.values(), key=lambda x: -x["usable_tb"])


def build(rows: list[dict], out_path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=13)
    wb = Workbook()

    def sheet(ws, columns, data, title):
        ws.cell(row=1, column=1, value=title).font = title_font
        for ci, (f, label, w) in enumerate(columns, start=1):
            c = ws.cell(row=3, column=ci, value=label)
            c.fill, c.font = header_fill, header_font
            c.alignment = Alignment(horizontal="center")
        for ri, row in enumerate(data, start=4):
            for ci, (f, _l, _w) in enumerate(columns, start=1):
                ws.cell(row=ri, column=ci, value=row.get(f))
        for ci, (_f, _l, w) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A4"

    ws1 = wb.active
    ws1.title = "Arrays"
    sheet(ws1, [
        ("array_name", "Array", 30), ("vendor", "Vendor", 12),
        ("cloud_type", "Cloud Type", 12), ("csp", "CSP", 12),
        ("group", "Group", 20), ("model", "Model", 24), ("status", "Status", 12),
        ("usable_tb", "Usable (TB)", 13), ("used_tb", "Used (TB)", 12),
        ("free_tb", "Free (TB)", 12), ("utilization_pct", "Util %", 9),
        ("data_reduction", "Data Reduction", 14), ("collected_at", "Collected At", 26),
    ], rows, f"USM Capacity — Per Array  ({len(rows)} arrays)")

    base = [("arrays", "Arrays", 10), ("usable_tb", "Usable (TB)", 13),
            ("used_tb", "Used (TB)", 12), ("free_tb", "Free (TB)", 12),
            ("utilization_pct", "Util %", 9)]

    sheet(wb.create_sheet("By Cloud Type"), [("cloud_type", "Cloud Type", 16)] + base,
          summarize(rows, "cloud_type"), "Capacity by Cloud Type (On-Prem vs Cloud)")
    sheet(wb.create_sheet("By CSP"), [("csp", "CSP", 16)] + base,
          summarize(rows, "csp"), "Capacity by Cloud Provider (CSP)")
    sheet(wb.create_sheet("By Vendor"), [("vendor", "Vendor", 16)] + base,
          summarize(rows, "vendor"), "Capacity by Platform (Vendor)")
    sheet(wb.create_sheet("By Group"), [("group", "Group", 24)] + base,
          summarize(rows, "group"), "Capacity by Group")

    pivot: dict = {}
    for r in rows:
        k = (r["vendor"], r["csp"])
        a = pivot.setdefault(k, {"vendor": r["vendor"], "csp": r["csp"],
                                 "arrays": 0, "usable_tb": 0.0, "used_tb": 0.0, "free_tb": 0.0})
        a["arrays"] += 1
        a["usable_tb"] += r["usable_tb"]
        a["used_tb"] += r["used_tb"]
        a["free_tb"] += r["free_tb"]
    for a in pivot.values():
        a["usable_tb"] = round(a["usable_tb"], 2)
        a["used_tb"] = round(a["used_tb"], 2)
        a["free_tb"] = round(a["free_tb"], 2)
        a["utilization_pct"] = round(a["used_tb"] / a["usable_tb"] * 100, 1) if a["usable_tb"] > 0 else 0.0
    sheet(wb.create_sheet("Vendor x CSP"),
          [("vendor", "Vendor", 16), ("csp", "CSP", 14)] + base,
          sorted(pivot.values(), key=lambda x: -x["usable_tb"]), "Capacity by Vendor × CSP")

    wb.save(out_path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base-url", default="http://usodclpsandadm1.corp.intranet:8000")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    out = args.out or f"usm_capacity_{dt.datetime.now():%Y%m%d_%H%M}.xlsx"
    print(f"Fetching arrays from {args.base_url} …")
    arrays = fetch_arrays(args.base_url)
    rows = to_rows(arrays)
    print(f"  {len(rows)} arrays. Building workbook → {out}")
    build(rows, out)
    print("Done.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
